"""
File reader tool.
Supports: PDF, TXT, CSV, DOCX, XLSX, JSON.
Files must have been previously uploaded and stored in the uploads/ directory.
"""

import csv
import io
import json
import os
from pathlib import Path

from backend.tools.base import BaseTool, ToolResult
from backend.utils.config import get_settings
from backend.utils.logger import get_logger

logger = get_logger(__name__)

MAX_TEXT_CHARS = 12_000  # cap to avoid blowing out the LLM context


class FileReaderTool(BaseTool):
    """Reads the content of an uploaded file."""

    name = "read_file"
    description = (
        "Reads and extracts text content from an uploaded file. "
        "Supported formats: PDF, TXT, CSV, DOCX, XLSX, JSON. "
        "Use the filename exactly as returned by the file upload endpoint."
    )

    async def execute(self, arguments: dict) -> ToolResult:
        filename: str = arguments.get("filename", "").strip()
        if not filename:
            return ToolResult(success=False, error="No filename provided.")

        settings = get_settings()
        upload_dir = Path(settings.upload_dir)
        file_path = upload_dir / filename

        # Security: prevent path traversal
        try:
            file_path = file_path.resolve()
            upload_dir_resolved = upload_dir.resolve()
            file_path.relative_to(upload_dir_resolved)
        except (ValueError, RuntimeError):
            return ToolResult(
                success=False,
                error="Invalid filename — path traversal detected.",
            )

        if not file_path.exists():
            return ToolResult(
                success=False,
                error=f"File '{filename}' not found in uploads directory.",
            )

        suffix = file_path.suffix.lower()
        logger.info("FileReader reading: %s (type: %s)", filename, suffix)

        try:
            if suffix == ".txt":
                return await self._read_text(file_path)
            elif suffix == ".json":
                return await self._read_json(file_path)
            elif suffix == ".csv":
                return await self._read_csv(file_path)
            elif suffix == ".pdf":
                return await self._read_pdf(file_path)
            elif suffix == ".docx":
                return await self._read_docx(file_path)
            elif suffix in (".xlsx", ".xls"):
                return await self._read_xlsx(file_path)
            else:
                return ToolResult(
                    success=False,
                    error=f"Unsupported file type: {suffix}.",
                )
        except Exception as exc:
            logger.error("FileReader error: %s", exc)
            return ToolResult(success=False, error=str(exc))

    # ── Readers ──────────────────────────────────────────────────────────

    async def _read_text(self, path: Path) -> ToolResult:
        text = path.read_text(encoding="utf-8", errors="replace")
        return ToolResult(
            success=True,
            result={"content": text[:MAX_TEXT_CHARS], "type": "txt"},
        )

    async def _read_json(self, path: Path) -> ToolResult:
        data = json.loads(path.read_text(encoding="utf-8"))
        content = json.dumps(data, indent=2)
        return ToolResult(
            success=True,
            result={"content": content[:MAX_TEXT_CHARS], "type": "json"},
        )

    async def _read_csv(self, path: Path) -> ToolResult:
        lines = []
        with open(path, newline="", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for i, row in enumerate(reader):
                lines.append(", ".join(row))
                if i > 200:  # limit rows in context
                    lines.append("... (truncated)")
                    break
        content = "\n".join(lines)
        return ToolResult(
            success=True,
            result={"content": content[:MAX_TEXT_CHARS], "type": "csv"},
        )

    async def _read_pdf(self, path: Path) -> ToolResult:
        try:
            import pymupdf as fitz  # PyMuPDF (fitz alias)
        except ImportError:
            return ToolResult(
                success=False,
                error="PyMuPDF (fitz) not installed. Run: pip install pymupdf",
            )
        doc = fitz.open(str(path))
        pages = []
        for i, page in enumerate(doc):
            pages.append(f"--- Page {i + 1} ---\n{page.get_text()}")
            if sum(len(p) for p in pages) > MAX_TEXT_CHARS:
                break
        content = "\n".join(pages)
        return ToolResult(
            success=True,
            result={
                "content": content[:MAX_TEXT_CHARS],
                "type": "pdf",
                "pages": doc.page_count,
            },
        )

    async def _read_docx(self, path: Path) -> ToolResult:
        try:
            from docx import Document
        except ImportError:
            return ToolResult(
                success=False,
                error="python-docx not installed. Run: pip install python-docx",
            )
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = "\n".join(paragraphs)
        return ToolResult(
            success=True,
            result={"content": content[:MAX_TEXT_CHARS], "type": "docx"},
        )

    async def _read_xlsx(self, path: Path) -> ToolResult:
        try:
            import openpyxl
        except ImportError:
            return ToolResult(
                success=False,
                error="openpyxl not installed. Run: pip install openpyxl",
            )
        wb = openpyxl.load_workbook(str(path), data_only=True)
        lines = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            lines.append(f"=== Sheet: {sheet} ===")
            for row in ws.iter_rows(values_only=True):
                row_str = ", ".join(str(c) if c is not None else "" for c in row)
                lines.append(row_str)
                if sum(len(l) for l in lines) > MAX_TEXT_CHARS:
                    lines.append("... (truncated)")
                    break
        content = "\n".join(lines)
        return ToolResult(
            success=True,
            result={"content": content[:MAX_TEXT_CHARS], "type": "xlsx"},
        )

    def schema(self) -> dict:
        return {
            "name": self.name,
            "description": self.description,
            "parameters": {
                "type": "object",
                "properties": {
                    "filename": {
                        "type": "string",
                        "description": (
                            "The exact filename of the uploaded file as returned "
                            "by the upload endpoint. "
                            "Example: 'report_abc123.pdf'."
                        ),
                    }
                },
                "required": ["filename"],
            },
        }


# Singleton instance
file_reader_tool = FileReaderTool()
